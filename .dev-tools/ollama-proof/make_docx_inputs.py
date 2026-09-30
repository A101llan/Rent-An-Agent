"""Generate the messy .docx meeting transcript used by the Ollama proofs (python-docx)."""

from pathlib import Path

import docx
from docx.shared import Pt

OUT = Path(__file__).parent / "inputs" / "meeting-notes" / "03-client-call-kariuki-logistics.docx"

LINES = [
    ("h", "Kariuki Logistics x AgentHub - pilot call (Otter export, unedited)"),
    ("p", "Recorded Wed 30 Sep 2026. Speakers auto-labelled, some labels wrong."),
    ("s", "Ge Wambua  00:00:04", "Hello hello, habari za asubuhi? Can you hear me? I think Lucy is still joining."),
    ("s", "Mzee Kariuki  00:00:11", "Nzuri sana. Yes we hear you. Lucy is coming, she's in traffic, you know Mombasa Road at this hour."),
    ("s", "Achieng Ouma  00:00:19", "Haha yes, eeh, Mombasa Road. Morning Mzee."),
    ("s", "Mzee Kariuki  00:00:24", "Morning morning. So, um, where do we start?"),
    ("s", "Ge Wambua  00:00:29", "Let me give a quick recap. Last time we talked about the meeting notes agent and the invoice checker running on your depot PCs, locally, so nothing goes to the cloud."),
    ("s", "Mzee Kariuki  00:00:41", "Yes, that was the main thing for us. Our board is very, very strict about data. Customer names, delivery addresses, all that, it cannot leave our building."),
    ("s", "Achieng Ouma  00:00:52", "Right, and that's exactly why we built the local runtime. It runs on the PC with a small model, Ollama, and the files stay on the machine."),
    ("s", "Lucy Wairimu  00:01:03", "Sorry sorry I'm here now. Traffic ilikuwa mbaya. Did I miss anything?"),
    ("s", "Ge Wambua  00:01:08", "No no, we just started. Karibu Lucy."),
    ("s", "Lucy Wairimu  00:01:12", "Thanks. So from procurement side, my questions are mostly about cost and the contract."),
    ("s", "Mzee Kariuki  00:01:20", "Let's do scope first. How many depots for the pilot? I was thinking all five. Nairobi, Mombasa, Kisumu, Nakuru, Eldoret."),
    ("s", "Ge Wambua  00:01:31", "Five is fine from our side. The installer is the same for all of them."),
    ("s", "Achieng Ouma  00:01:36", "Yeah five works, we'd need one IT contact per depot though."),
    ("s", "Speaker 5  00:01:44", "[crosstalk] ...the Eldoret one has the old Dell..."),
    ("s", "Mzee Kariuki  00:01:47", "Sorry who was that? Ah, that's Otieno from IT, he's dialled in from Eldoret. Otieno, go ahead."),
    ("s", "Otieno (IT)  00:01:53", "Yes sorry, I was saying the Eldoret depot PC is an old Dell, maybe 4 gigs RAM. I don't think it can run a model."),
    ("s", "Ge Wambua  00:02:01", "Ok good to know. The 1b model wants about 8 gigs to be comfortable."),
    ("s", "Lucy Wairimu  00:02:09", "And budget-wise, honestly, five depots in the pilot is too much for this quarter. Can we start smaller?"),
    ("s", "Mzee Kariuki  00:02:16", "Hmm. Ok. Lucy is right, I've seen the numbers. Let's do three depots. Nairobi, Mombasa and Nakuru. Kisumu and Eldoret in phase two."),
    ("s", "Ge Wambua  00:02:27", "Three depots, Nairobi, Mombasa, Nakuru. Noted. That's the pilot scope then?"),
    ("s", "Mzee Kariuki  00:02:31", "Yes, final. Three."),
    ("s", "Achieng Ouma  00:02:35", "Perfect. That actually makes the rollout easier for us too."),
    ("s", "Lucy Wairimu  00:02:40", "Now pricing. What does the pilot cost us?"),
    ("s", "Ge Wambua  00:02:44", "For the pilot we were proposing thirty days free, then per-run pricing after, KES 50 per run like our standard local plan."),
    ("s", "Lucy Wairimu  00:02:53", "Thirty days free, ok, I like that. And after thirty days we can cancel without penalty?"),
    ("s", "Ge Wambua  00:02:59", "Yes, no penalty, you can cancel anytime."),
    ("s", "Lucy Wairimu  00:03:03", "Then I'm fine with that. Mzee?"),
    ("s", "Mzee Kariuki  00:03:06", "Agreed. Thirty days free, then KES 50 per run."),
    ("s", "Achieng Ouma  00:03:12", "Great. So I'll send the pilot agreement draft, I can have it to you by Monday."),
    ("s", "Lucy Wairimu  00:03:18", "Monday is good. Send it to me directly, not the generic procurement inbox, that one is a black hole."),
    ("s", "Achieng Ouma  00:03:24", "Haha noted, directly to you."),
    ("s", "Mzee Kariuki  00:03:28", "Also, eh, one thing. The drivers send voice notes in Swahili on WhatsApp about deliveries. Can the notes thing handle voice notes?"),
    ("s", "Ge Wambua  00:03:38", "Not today. It's text files, markdown, Word documents. Voice would need a speech model on the PC, we haven't tested that on these machines."),
    ("s", "Mzee Kariuki  00:03:47", "Ok, but will the pilot support Swahili voice notes at some point? That's what my depot managers will ask me."),
    ("s", "Ge Wambua  00:03:54", "Honestly I don't know yet. Let us look into it, I can't promise for the pilot."),
    ("s", "Mzee Kariuki  00:03:59", "Fine, fine. Let's keep it as a question."),
    ("s", "Lucy Wairimu  00:04:03", "Another one from me. If the PCs are too weak, like Otieno said about Eldoret, who pays for the RAM upgrades? Us or you?"),
    ("s", "Ge Wambua  00:04:11", "Hmm. Good question. We don't normally supply hardware..."),
    ("s", "Mzee Kariuki  00:04:15", "Let's not decide today. Eldoret is phase two anyway."),
    ("s", "Lucy Wairimu  00:04:19", "Ok but it will come up. Write it down."),
    ("s", "Otieno (IT)  00:04:23", "Also before anything gets installed I need the security team to sign off on Ollama. They'll want to know what ports it opens."),
    ("s", "Achieng Ouma  00:04:31", "It only listens on 127.0.0.1, port 11434, nothing exposed to the network. We can send a one-pager."),
    ("s", "Mzee Kariuki  00:04:39", "Good. I will get the IT security sign-off myself, I'll talk to the CISO this week. Otieno, you just support."),
    ("s", "Otieno (IT)  00:04:46", "Sawa Mzee."),
    ("s", "Lucy Wairimu  00:04:48", "And I'll share the list of the three depots with the contact person for each, by Wednesday."),
    ("s", "Ge Wambua  00:04:54", "Perfect, thanks Lucy."),
    ("s", "Achieng Ouma  00:04:57", "Someone also needs to check whether the Nairobi, Mombasa and Nakuru PCs actually have 8 gigs of RAM. Before we ship the installer."),
    ("s", "Mzee Kariuki  00:05:04", "Yes that should be checked."),
    ("s", "Ge Wambua  00:05:06", "Mm, yeah. Definitely needs doing."),
    ("s", "Speaker 5  00:05:09", "[inaudible]"),
    ("s", "Lucy Wairimu  00:05:12", "Can I also ask, does this integrate with SAP? Our invoices are all in SAP Business One."),
    ("s", "Achieng Ouma  00:05:19", "Not directly for the pilot. You'd export invoices as files. A proper SAP connector we'd have to scope."),
    ("s", "Lucy Wairimu  00:05:26", "Ok so SAP integration is open. Fine for the pilot, but I'll need an answer before we sign the full contract."),
    ("s", "Mzee Kariuki  00:05:34", "Anything else? I have another call at half past."),
    ("s", "Ge Wambua  00:05:38", "One small thing, the model runs fully offline. We decided earlier that no data goes to any cloud LLM, not even as a backup. Is that still the requirement?"),
    ("s", "Mzee Kariuki  00:05:47", "Yes. No cloud fallback. If the local thing fails it should just fail and tell us, not send data anywhere."),
    ("s", "Ge Wambua  00:05:54", "Understood, that's how it works."),
    ("s", "Lucy Wairimu  00:05:57", "Ok. Asanteni sana."),
    ("s", "Achieng Ouma  00:05:59", "Asante, talk soon. Have a good day Mzee."),
    ("s", "Mzee Kariuki  00:06:02", "Kwaheri."),
]

# Pad with realistic small talk / off-topic chatter so the file exceeds one 6000-char chunk.
SMALLTALK = [
    ("s", "Achieng Ouma  00:00:33", "Before we go on, Mzee, how was the Naivasha trip? I saw the photos on LinkedIn, the flamingos!"),
    ("s", "Mzee Kariuki  00:00:36", "Ah it was good, too short. The kids wanted to stay. Next time we go for a full week, maybe December, eh."),
    ("s", "Ge Wambua  00:00:38", "December in Naivasha is nice but the traffic on the escarpment, wueh. Last year it took us six hours."),
    ("s", "Mzee Kariuki  00:00:39", "Six hours! No, you must leave at 5am. Anyway, anyway. Let's continue."),
]

SMALLTALK2 = [
    ("s", "Otieno (IT)  00:03:13", "Sorry, side note, did you guys watch Harambee Stars on Sunday? That second goal, offside kabisa."),
    ("s", "Achieng Ouma  00:03:15", "Eeh, I was at Nyayo! The referee was blind, the whole stadium was shouting."),
    ("s", "Mzee Kariuki  00:03:17", "Otieno, focus. But yes, it was offside. Haha. Ok."),
    ("s", "Lucy Wairimu  00:03:19", "Also my M-Pesa was down for an hour on Saturday, I couldn't even pay for parking at the Junction. The askari was not amused."),
    ("s", "Ge Wambua  00:03:21", "Same here, I think the whole network had issues. Anyway, back to the pilot, sorry."),
    ("s", "Otieno (IT)  00:03:22", "One more random thing, the Eldoret generator failed again last week, we lost power for two days. So even phase two, we need to think about power there."),
    ("s", "Mzee Kariuki  00:03:24", "That's a separate facilities issue, not for this call. Let's move."),
    ("s", "Achieng Ouma  00:03:25", "Sure. For context, other clients in the pilot program have been using it for board meeting minutes and supplier calls, mostly Word documents, some markdown from Notion exports."),
    ("s", "Lucy Wairimu  00:03:26", "Notion, eh, we use OneNote mostly. Can we export OneNote to Word? I think yes."),
    ("s", "Ge Wambua  00:03:27", "Yes, OneNote exports to .docx fine. That works today."),
]
def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    i = next(n for n, l in enumerate(LINES) if l[1].startswith("Mzee Kariuki  00:03:28"))
    lines = LINES[:7] + SMALLTALK + LINES[7:i] + SMALLTALK2 + LINES[i:]
    for kind, *rest in lines:
        if kind == "h":
            d.add_heading(rest[0], level=1)
        elif kind == "p":
            d.add_paragraph(rest[0]).italic = True
        else:
            who, text = rest
            para = d.add_paragraph()
            para.add_run(who).bold = True
            para.add_run("\n" + text)
    d.save(OUT)
    text = "\n".join(p.text for p in docx.Document(str(OUT)).paragraphs)
    print(f"wrote {OUT} ({len(text)} chars when read back)")


if __name__ == "__main__":
    main()