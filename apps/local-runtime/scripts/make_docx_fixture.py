"""Generate fixtures/sample-meeting.docx (python-docx). Re-run to regenerate."""

from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
out = ROOT / "fixtures" / "sample-meeting.docx"

d = docx.Document()
d.add_heading("Mobile Checkout Sync - 29 Sep 2026", level=1)
d.add_paragraph("Attendees: Amina (PM), Brian (Backend), Carol (Design), David (QA)")
d.add_heading("Discussion", level=2)
d.add_paragraph(
    "Amina opened with the M-Pesa checkout drop-off: 18% of users abandon at the STK push step. "
    "Brian said callback latency from the payments gateway averages 9 seconds and sometimes times out."
)
d.add_paragraph(
    "Carol showed a new waiting screen with a countdown and a 'resend prompt' button. "
    "The team agreed to ship the new waiting screen in the 2.4 release."
)
d.add_paragraph(
    "David reported the Android 10 crash on the receipt page is still reproducible. "
    "Decision: the receipt-page crash blocks the 2.4 release until fixed."
)
d.add_heading("Action items", level=2)
t = d.add_table(rows=1, cols=3)
t.rows[0].cells[0].text = "Owner"
t.rows[0].cells[1].text = "Task"
t.rows[0].cells[2].text = "Due"
for owner, task, due in [
    ("Brian", "Add retry with idempotency key to the STK push callback handler", "Oct 3"),
    ("Carol", "Hand off final waiting-screen designs to mobile team", "Oct 1"),
    ("David", "Write regression test for the Android 10 receipt crash", "Oct 2"),
]:
    row = t.add_row().cells
    row[0].text, row[1].text, row[2].text = owner, task, due
d.add_heading("Open questions", level=2)
d.add_paragraph("Should we fall back to card payments automatically after two failed STK pushes?")
d.add_paragraph("Who owns the payments-gateway SLA conversation with the provider?")
out.parent.mkdir(parents=True, exist_ok=True)
d.save(out)
print(f"wrote {out}")
